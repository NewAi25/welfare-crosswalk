"""BUILD OUTPUT: turn the research tables into the one Excel file that answers the question.

In plain words
    The research is recorded in four tables in 2_research/. This script reads them,
    adds up the ease of adoption score exactly as 2_research/rules.md says, and writes
    4_output/welfare_crosswalk_dairy.xlsx. Nothing in the Excel file is typed by hand:
    every cell comes from a table row or from the Stygar product list, so the Excel file
    and the tables can never disagree.

The sheets it writes
    Read me             the question, how to read the file, the answers so far, how to check
    Answer              one row per welfare indicator, ranked, with every factor visible
    Sensors x products  every indicator against all 129 products in the Stygar list
    Schemes             every indicator against the four certification schemes
    Product list        the 129 products, with what each vendor says it does
    Sources             every document used, with its link

The score, in plain words (the full rule is in 2_research/rules.md, "The score")
    Five factors, each turned into 0 to 3 points, then added (0 to 15):
    sensor evidence      validated product 3, product without validation 2, research only 1, otherwise 0
    device needed        none extra 3, one per farm 2, one per animal 1, otherwise 0
    vendors naming it    10 or more 3, 3 to 9 2, 1 or 2 1, none 0
    named by FARM        with a number 3, named 2, only a related housing rule 1, otherwise 0
    schemes requiring it all four 3, two or three 2, one 1, none 0
    The score is a sort order, not a welfare judgement.

How to run it
    python 3_code/build_output.py           write the Excel file
    python 3_code/build_output.py --check   confirm the Excel file matches the tables (used by the push gate)
"""
import csv
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import count_vendors  # noqa: E402  (same folder)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TABLES = ROOT / "2_research"
OUTPUT = ROOT / "4_output" / "welfare_crosswalk_dairy.xlsx"
SCHEMES = ["RSPCA_Assured", "FARM_v5", "GAP_dairy", "Certified_Humane"]
SCHEME_NAMES = {"RSPCA_Assured": "RSPCA Assured (UK)", "FARM_v5": "FARM v5 (US industry)",
                "GAP_dairy": "Global Animal Partnership (US)", "Certified_Humane": "Certified Humane (US)"}

# The 0 to 3 points for each factor. These five lines are the only place the ranking is set,
# and they must match the table "The score" in 2_research/rules.md word for word.
SENSOR_POINTS = {"validated_commercial": 3, "commercial_unvalidated": 2, "research_only": 1}
DEVICE_POINTS = {"routine_data": 3, "farm_fixed": 2, "animal_mounted": 1}
FARM_POINTS = {"named_with_threshold": 3, "named": 2, "related_resource": 1}


def vendor_points(n):
    return 3 if n >= 10 else 2 if n >= 3 else 1 if n >= 1 else 0


def scheme_points(n):
    return 3 if n == 4 else 2 if n >= 2 else 1 if n == 1 else 0


# Plain English for every code, shown in the Excel file. Same meanings as 2_research/rules.md.
SENSOR_LABEL = {
    "validated_commercial": "Tested: a product in the Stygar list was tested for this in a published study; "
                            "the column 'What was tested, and how well' says how it did",
    "commercial_unvalidated": "Claimed: a product's own description names it; Stygar 2021 coded no validation study",
    "research_only": "Research only: no product in the Stygar list claims this measure; researchers propose a "
                     "method or a substitute measure",
    "none": "Nothing found: neither sensor source names a product or a method for it",
    "unsure": "Unclear: the written rules do not settle it (see open questions)",
}
DEVICE_LABEL = {
    "routine_data": "Nothing extra: already in milk records or farm records",
    "farm_fixed": "One device per farm or barn (camera, scale, trough meter)",
    "animal_mounted": "One device per animal (collar, ear tag, bolus)",
    "sample_or_procedure": "A sample or a scoring by a person",
    "none": "No technology named in the sources",
    "unsure": "Unclear",
}
FARM_LABEL = {
    "named_with_threshold": "Named in FARM, with a number to meet",
    "named": "Named in FARM",
    "related_resource": "FARM has only a related housing, equipment or practice rule",
    "not_named": "Not named in FARM",
    "unsure": "Unclear whether FARM names it",
}
HOW_CHECKED_LABEL = {
    "visual_inspection": "inspector looks at the animals",
    "records_review": "inspector checks farm records",
    "sensor_accepted": "sensor data accepted",
    "unspecified": "required, but the scheme does not say how it is checked",
    "unsure": "unclear from the scheme's text",
}
MANUAL = ("visual_inspection", "records_review")
MEASURE_TYPE_LABEL = {"animal_based": "Observed or recorded on the animal", "resource_based": "About the housing or equipment",
                      "management_based": "About farm management", "none": "No measure defined"}
LISTED_BY_LABEL = {"joined": "EFSA and Welfare Quality", "efsa_only": "EFSA only", "wq_only": "Welfare Quality only"}
NOT_YET = "Not yet checked"

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
VALIDATED_FILL = PatternFill("solid", fgColor="70AD47")
NAMED_FILL = PatternFill("solid", fgColor="E2EFDA")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")


def read(name):
    with (TABLES / name).open(newline="", encoding="utf-8-sig") as f:
        return [r for r in csv.DictReader(f) if any((v or "").strip() for v in r.values())]


def indicator_name(r):
    if r["efsa_measure"]:
        return f"{r['efsa_welfare_consequence']}: {r['efsa_measure']}"
    return r["wq_measure"] or f"{r['wq_criterion']} (Welfare Quality defines no measure)"


def compute():
    """Everything the Excel file shows, as plain Python data. Used for writing and for --check."""
    indicators = read("indicators.csv")
    sensors = {r["indicator_id"]: r for r in read("sensor_coverage.csv")}
    schemes = {}
    for r in read("crosswalk.csv"):
        schemes.setdefault(r["indicator_id"], {})[r["scheme"]] = r
    products = count_vendors.appendix_products()
    profiled = {r["appendix_name"]: r["product_id"] for r in read("products.csv") if r["appendix_name"]}

    answer = []
    for ind in indicators:
        iid = ind["indicator_id"]
        s = sensors.get(iid)
        rows = schemes.get(iid, {})
        schemes_done = all(k in rows for k in SCHEMES)
        required = [k for k in SCHEMES if rows.get(k, {}).get("requirement_quote", "").strip()]
        grade = s["sensor_evidence"] if s else ""
        vendors = int(s["vendors_naming_it"]) if s else 0
        farm = rows.get("FARM_v5", {}).get("farm_naming", "") or ("not_named" if "FARM_v5" in rows else "")
        points = [SENSOR_POINTS.get(grade, 0), DEVICE_POINTS.get(s["device_needed"], 0) if s else 0,
                  vendor_points(vendors), FARM_POINTS.get(farm, 0), scheme_points(len(required))]

        # Kevin's list 1: a sensor measures it, and no scheme requires it.
        if not s:
            list1 = NOT_YET
        elif grade == "unsure":
            list1 = "Unclear: sensor evidence unsure (see open questions)"
        elif grade not in ("validated_commercial", "commercial_unvalidated"):
            list1 = "No"
        elif not schemes_done:
            list1 = "Not yet known: schemes not yet checked"
        else:
            list1 = "Yes" if not required else "No"
        # Kevin's list 2: a scheme requires it and checks it by hand (inspection or records), not by sensor data.
        manual = [k for k in required if rows[k]["how_checked"] in MANUAL]
        unclear = [k for k in required if rows[k]["how_checked"] == "unsure"]
        if manual:
            list2 = "Yes: " + ", ".join(SCHEME_NAMES[k] for k in manual)
        elif unclear:
            list2 = "Unclear: how " + ", ".join(SCHEME_NAMES[k] for k in unclear) + " checks it (see open questions)"
        elif not schemes_done:
            list2 = "Not yet known: schemes not yet checked"
        else:
            list2 = "No"

        status = ("Complete" if s and schemes_done else
                  "Sensors done, schemes not yet checked" if s else NOT_YET)
        answer.append({
            "id": iid, "name": indicator_name(ind), "listed_by": LISTED_BY_LABEL[ind["listed_by"]],
            "wq_name": ind["wq_measure"] if ind["listed_by"] == "joined" else "",
            "measure_type": MEASURE_TYPE_LABEL[ind["measure_type"]],
            "sensor": SENSOR_LABEL[grade] if s else NOT_YET,
            "tested_how": s["validated_trait"] if s and grade == "validated_commercial" else "",
            "validated": "; ".join(k.split("|", 1)[1] for k in filter(None, (x.strip() for x in s["validated_products"].split(";")))) if s else "",
            "vendors": vendors if s else NOT_YET,
            "device": DEVICE_LABEL[s["device_needed"]] if s else NOT_YET,
            "farm": FARM_LABEL.get(farm, NOT_YET),
            "schemes_n": len(required) if schemes_done else NOT_YET,
            "schemes_how": "; ".join(f"{SCHEME_NAMES[k]}: {HOW_CHECKED_LABEL[rows[k]['how_checked']]}" for k in required) if rows else NOT_YET,
            "list1": list1, "list2": list2,
            "score": sum(points), "parts": " + ".join(str(p) for p in points), "status": status,
            "where": s["where_to_check"] if s else "", "reasoning": s["reasoning"] if s else "",
            "grade": grade,
        })
    answer.sort(key=lambda r: (-r["score"], r["id"]))
    for rank, r in enumerate(answer, start=1):
        r["rank"] = rank

    matrix = []
    for ind in indicators:
        s = sensors.get(ind["indicator_id"])
        validated = {k.strip() for k in (s["validated_products"].split(";") if s else []) if k.strip()}
        words = count_vendors.search_words(s["search_words"]) if s else []
        cells = []
        for p in products:
            cells.append("validated" if p["key"] in validated else "named" if s and count_vendors.names_it(p, words) else "")
        matrix.append({"id": ind["indicator_id"], "name": indicator_name(ind),
                       "sensor": SENSOR_LABEL[s["sensor_evidence"]] if s else NOT_YET, "cells": cells, "done": bool(s)})

    scheme_rows = []
    for ind in indicators:
        rows = schemes.get(ind["indicator_id"], {})
        line = {"id": ind["indicator_id"], "name": indicator_name(ind), "per_scheme": []}
        for k in SCHEMES:
            r = rows.get(k)
            if r is None:
                line["per_scheme"].append((NOT_YET, "", ""))
            elif r["requirement_quote"].strip():
                where = f"{r['section']}, page {r['page']}"
                line["per_scheme"].append(("Required", HOW_CHECKED_LABEL[r["how_checked"]], f"\"{r['requirement_quote']}\" ({where})"))
            else:
                line["per_scheme"].append(("Not required", "", r["reasoning"]))
        scheme_rows.append(line)

    with (ROOT / "1_sources" / "manifest.csv").open(newline="", encoding="utf-8") as f:
        sources = list(csv.DictReader(f))
    return {"indicators": indicators, "answer": answer, "matrix": matrix, "schemes": scheme_rows,
            "products": products, "profiled": profiled, "sources": sources, "sensors": sensors}


def summary_lines(d):
    a = d["answer"]
    done = [r for r in a if r["grade"]]
    grades = Counter(r["grade"] for r in done)
    lines = [
        ("The welfare indicators", f"{len(a)} indicators: "
         + ", ".join(f"{n} {LISTED_BY_LABEL[k]}" for k, n in sorted(Counter(i['listed_by'] for i in d['indicators']).items()))),
        ("(a) Can a sensor measure it?", f"Checked for {len(done)} of {len(a)} indicators so far. "
         + "; ".join(f"{grades.get(k, 0)} {SENSOR_LABEL[k].split(':')[0].lower()}" for k in SENSOR_LABEL)),
        ("(b) How easy to adopt?", "See the Answer sheet: device needed, products whose description names it, and whether FARM (the US industry "
         "programme) names it. " + f"FARM is checked for {sum(1 for r in a if r['farm'] != NOT_YET)} of {len(a)} indicators so far."),
        ("(c) Already covered by schemes?", f"Checked for {sum(1 for r in a if r['schemes_n'] != NOT_YET)} of {len(a)} indicators so far."),
        ("Kevin's list 1: Tested or Claimed in column (a), required by no scheme",
         f"{sum(1 for r in a if r['list1'] == 'Yes')} confirmed; {sum(1 for r in a if r['list1'].startswith('Not yet known'))} "
         f"are Tested or Claimed and wait for the scheme check; {sum(1 for r in a if r['list1'].startswith('Unclear'))} more are "
         "unclear because their sensor evidence is unsure."),
        ("Kevin's list 2: required by a scheme, checked by hand",
         f"{sum(1 for r in a if r['list2'].startswith('Yes'))} confirmed; the schemes are checked for "
         f"{sum(1 for r in a if r['schemes_n'] != NOT_YET)} of {len(a)} indicators so far. Column (a) shows the sensor "
         "evidence beside each one."),
    ]
    return lines


def write(d, path=OUTPUT):
    wb = Workbook()

    # Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 110
    text = [
        ("Welfare crosswalk, dairy cows", ""),
        ("The question", "Kevin Xia's Top 30 list, item 13, asks for a crosswalk of welfare indicators against what "
         "commercial sensors measure and what certification schemes require and audit, for one species. The meeting "
         "tool's notes of a call on 22 September 2026 record him proposing to narrow it to indicators that are (a) measurable by AI, (b) acceptable to "
         "industry and (c) covered by certification schemes; the written scope that followed also counted coverage by "
         "existing sensors. On 25 September he called industry acceptability and both kinds of coverage soft "
         "requirements that make adoption easier, not hard ones. The three files are in 1_sources (kevin_*.md)."),
        ("The two lists Kevin asked for", "List 1: indicators graded Tested or Claimed in column (a) that no scheme requires. "
         "List 2: indicators a scheme requires but checks by hand (an inspector looks at the animals or the farm "
         "records). Column (a) shows the sensor evidence beside each indicator on either list."),
        ("How to read this file", "Answer: one row per indicator, sorted by the ease of adoption score. The score is a "
         "rule of thumb used only as a sort order, not a recommendation; every factor has its own column, so you can "
         "sort by any of them. Sensors x products: for each indicator, which of the 129 products in the Stygar 2021 list "
         "name it in their own description and which a published study tested. Schemes: what each certification scheme "
         "requires. Product list and Sources: where everything comes from."),
        ("Status", "Work in progress. Rows marked 'Not yet checked' have no sensor or scheme research yet; their score is not "
         "meaningful and the ranking is provisional until every row is complete."),
        ("", ""),
        ("THE ANSWERS SO FAR", ""),
    ] + summary_lines(d) + [
        ("", ""),
        ("WHAT THE LABELS MEAN", ""),
    ] + [(f"Sensor: {k}", v) for k, v in SENSOR_LABEL.items()] + [
        ("Device needed", "The cheapest kind of equipment that measures it. Classes, in the order the score ranks them: "
         "nothing extra (3 points), one device per farm (2), one device per animal (1), a sample or a scoring by a "
         "person (0), or no technology named in the sources (0). 'Nothing extra' can appear where no sensor exists, "
         "because the value is already taken from milk or farm records (2_research/rules.md, rule 4.4)."),
        ("Vendors naming it", "How many of the 129 products in the Stygar list contain one of the indicator's search "
         "words in their own description. A count of claims, not of proof. 'Tested products' is a separate column: a "
         "product can be tested without its description naming the measure, and the reverse."),
        ("Sensors x products: 'validated'", "A published study tested this product for this measure (Stygar 2021, Tables 1 and 2)."),
        ("Sensors x products: 'named'", "The vendor's own description contains one of the indicator's search words. Not tested."),
        ("Ease of adoption score", "0 to 15, five factors of 0 to 3 added together (2_research/rules.md, section 5). "
         "Ties are sorted by indicator id. "
         "A sort order, not a welfare judgement. Score parts are shown in the order: sensor + device + vendors + FARM + schemes."),
        ("", ""),
        ("WHAT THIS DOES NOT SHOW", "The sensor evidence is a baseline from two papers, Stygar et al. 2021 and Maroto "
         "Molina et al. 2020. Products, studies and links are as those papers found them; 'Nothing found' means nothing "
         "in those two papers, not nothing anywhere."),
        ("", "'Tested' means a published study tested the product for the measure, as Stygar 2021 coded it; many of "
         "those tests fell below Stygar's high performance bar, which the column 'What was tested, and how well' shows."),
        ("", "The certification schemes have not been read yet, so Kevin's two lists are not yet known."),
        ("", "The score is a rule of thumb used as a sort order. It is not a welfare judgement, and no welfare scientist "
         "has reviewed the table yet."),
        ("", ""),
        ("HOW TO CHECK ANY CELL", "Every checked row of the Answer sheet has a 'Where to check' column naming the "
         "documents, with pages where a passage is cited. The documents are listed in the Sources sheet and in "
         "1_sources/. The rules that turn what a page says into a label are in 2_research/rules.md, and section 2 of "
         "that file explains the codes, abbreviations and terms used in the cells. "
         "2_research/README.md explains the three kinds of check: read the page, apply the rule, or count again."),
        ("Who checked it", "Before each push, an AI auditing agent checks the changes against the source pages; "
         "2_research/audit_log.md lists the versions that passed. The author plans to spot check ten rows after each "
         "push; each spot check is logged in 2_research/decision_log.md with its date and the rows read."),
        ("Built from", "The tables in 2_research/ by 3_code/build_output.py. Do not edit this file by hand; change "
         "the tables and run the script."),
    ]
    for row, (a, b) in enumerate(text, start=1):
        ws.cell(row=row, column=1, value=a).font = Font(bold=True, size=14 if row == 1 else 11)
        ws.cell(row=row, column=2, value=b).alignment = WRAP

    # Answer
    ws = wb.create_sheet("Answer")
    columns = [
        ("Rank", 6, "rank"), ("ID", 7, "id"), ("Welfare indicator", 34, "name"), ("Listed by", 16, "listed_by"),
        ("Welfare Quality name, if different", 22, "wq_name"), ("What is measured", 16, "measure_type"),
        ("(a) Can a sensor measure it?", 34, "sensor"), ("Tested products", 30, "validated"),
        ("What was tested, and how well (Stygar 2021, Table 2)", 34, "tested_how"),
        ("Products whose description names it (of 129)", 14, "vendors"),
        ("(b) Device needed", 28, "device"), ("(b) Named by FARM, the US industry programme", 22, "farm"),
        ("(c) Schemes requiring it (of 4)", 12, "schemes_n"), ("(c) Which schemes, and how they check", 30, "schemes_how"),
        ("Kevin list 1: Tested or Claimed in column (a), and no scheme requires it", 22, "list1"),
        ("Kevin list 2: required by a scheme, checked by hand (inspection or records)", 22, "list2"),
        ("Ease of adoption score (0 to 15, sort order only)", 12, "score"),
        ("Score parts (sensor + device + vendors + FARM + schemes)", 16, "parts"),
        ("Row status", 20, "status"), ("Where to check", 70, "where"), ("Reasoning", 60, "reasoning"),
    ]
    incomplete = sum(1 for r in d["answer"] if r["status"] != "Complete")
    if incomplete:
        columns[0] = (f"Rank (provisional: {incomplete} of {len(d['answer'])} rows incomplete)", 12, "rank")
    for c, (title, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    for row, r in enumerate(d["answer"], start=2):
        for c, (_, _, key) in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=c, value=r[key])
            cell.alignment = WRAP
            if r["status"] == NOT_YET:
                cell.fill = GREY_FILL
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(d['answer']) + 1}"

    # Sensors x products
    ws = wb.create_sheet("Sensors x products")
    ws.cell(row=1, column=1, value="Green 'validated': a published study tested this product for this measure. "
            "Light 'named': the vendor's description contains one of the row's search words; not tested. "
            "Blank: neither. Grey rows: not yet checked.")
    fixed = [("ID", 7), ("Welfare indicator", 30), ("(a) Can a sensor measure it?", 30)]
    for c, (title, width) in enumerate(fixed, start=1):
        cell = ws.cell(row=2, column=c, value=title)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    for i, p in enumerate(d["products"]):
        c = len(fixed) + 1 + i
        cell = ws.cell(row=2, column=c, value=f"{p['name']} ({p['provider']})")
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(text_rotation=90, vertical="bottom")
        ws.column_dimensions[get_column_letter(c)].width = 4
    ws.row_dimensions[2].height = 230
    for row, m in enumerate(d["matrix"], start=3):
        for c, value in enumerate([m["id"], m["name"], m["sensor"]], start=1):
            ws.cell(row=row, column=c, value=value).alignment = WRAP
        for i, v in enumerate(m["cells"]):
            cell = ws.cell(row=row, column=len(fixed) + 1 + i, value=v or None)
            if v == "validated":
                cell.fill = VALIDATED_FILL
            elif v == "named":
                cell.fill = NAMED_FILL
            elif not m["done"]:
                cell.fill = GREY_FILL
    ws.freeze_panes = "D3"

    # Schemes
    ws = wb.create_sheet("Schemes")
    head = [("ID", 7), ("Welfare indicator", 30)]
    for k in SCHEMES:
        head += [(f"{SCHEME_NAMES[k]}: required?", 14), (f"{SCHEME_NAMES[k]}: how checked", 18),
                 (f"{SCHEME_NAMES[k]}: quote and page, or why not", 40)]
    for c, (title, width) in enumerate(head, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    for row, line in enumerate(d["schemes"], start=2):
        values = [line["id"], line["name"]] + [v for triple in line["per_scheme"] for v in triple]
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v).alignment = WRAP
    ws.freeze_panes = "C2"

    # Product list
    ws = wb.create_sheet("Product list")
    head = [("Spreadsheet row", 8), ("Product", 30), ("Company", 22), ("Sensor type", 24),
            ("What the vendor says it does (Stygar list, column AIM)", 60), ("Country", 12),
            ("Validation study found by Stygar", 12), ("Profiled in products.csv", 10), ("Link (2021)", 40)]
    for c, (title, width) in enumerate(head, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    for row, p in enumerate(d["products"], start=2):
        values = [p["row"], p["name"], p["provider"], p["sensor_type"], p["aim"], p["country"], p["has_study"],
                  d["profiled"].get(p["name"], ""), p["link"]]
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v).alignment = WRAP
    ws.freeze_panes = "C2"

    # Sources
    ws = wb.create_sheet("Sources")
    keys = [k for k in ("title", "filename", "used_for", "licence", "in_repository", "doi_or_landing", "url", "retrieved")
            if k in d["sources"][0]]
    widths = {"title": 50, "filename": 30, "used_for": 40, "licence": 18, "in_repository": 10,
              "doi_or_landing": 40, "url": 40, "retrieved": 12}
    titles = {"title": "Document", "filename": "File in 1_sources", "used_for": "Used for", "licence": "Licence",
              "in_repository": "File in the repository?", "doi_or_landing": "Where to find it",
              "url": "Exact file the research read", "retrieved": "Downloaded"}
    for c, k in enumerate(keys, start=1):
        cell = ws.cell(row=1, column=c, value=titles[k])
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = widths[k]
    for row, s in enumerate(d["sources"], start=2):
        for c, k in enumerate(keys, start=1):
            ws.cell(row=row, column=c, value=s[k]).alignment = WRAP

    wb.properties.creator = "3_code/build_output.py"
    path.parent.mkdir(exist_ok=True)
    wb.save(path)


def check(d):
    """Rebuild in memory and compare every cell value with the saved Excel file."""
    if not OUTPUT.exists():
        print(f"{OUTPUT.name} is missing; run python 3_code/build_output.py")
        return 1
    import tempfile
    saved = load_workbook(OUTPUT)
    with tempfile.TemporaryDirectory() as tmp:
        fresh_path = Path(tmp) / "fresh.xlsx"
        write(d, fresh_path)
        fresh = load_workbook(fresh_path)
        for name in fresh.sheetnames:
            if name not in saved.sheetnames:
                print(f"sheet {name!r} missing from the saved file; run python 3_code/build_output.py")
                return 1
            a, b = fresh[name], saved[name]
            if a.max_row != b.max_row or a.max_column != b.max_column:
                print(f"sheet {name!r} has a different size; run python 3_code/build_output.py")
                return 1
            for ra, rb in zip(a.iter_rows(values_only=True), b.iter_rows(values_only=True)):
                if ra != rb:
                    print(f"sheet {name!r} is out of date; run python 3_code/build_output.py")
                    return 1
    print(f"{OUTPUT.name}: matches the tables")
    return 0


def main(argv):
    d = compute()
    if "--check" in argv:
        return check(d)
    write(d)
    a = d["answer"]
    print(f"wrote {OUTPUT.relative_to(ROOT)}")
    print(f"indicators: {len(a)}; sensor evidence filled for {sum(1 for r in a if r['grade'])}; "
          f"schemes checked for {sum(1 for r in a if r['schemes_n'] != NOT_YET)}")
    for label, text in summary_lines(d):
        print(f"  {label}: {text}")
    print("the ease of adoption score is a sort order, not a welfare judgement; see 2_research/rules.md")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
